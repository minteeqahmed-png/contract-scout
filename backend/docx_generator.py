import io
from docx import Document
from docx.shared import RGBColor, Pt
from .auditor import ContractAuditReport

def generate_redline_docx(report: ContractAuditReport) -> bytes:
    doc = Document()
    
    # Header
    doc.add_heading(f"Contract Audit Report: {report.contract_title}", 0)
    
    # Executive Summary
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(f"Total Contract Value: {report.total_contract_value}")
    
    p = doc.add_paragraph("Total Financial Exposure: ")
    run = p.add_run(report.total_financial_exposure)
    run.bold = True
    if report.has_unlimited_exposure:
        run.font.color.rgb = RGBColor(220, 38, 38)
        
    doc.add_paragraph(f"Vendor Bias Score: {report.overall_bias_score}/100")
    
    doc.add_heading("Rule Violations & Redlines", level=1)
    
    for result in report.results:
        if result.status in ("FAIL", "RISK"):
            doc.add_heading(f"{result.rule_name} ({result.status})", level=2)
            doc.add_paragraph(f"Finding: {result.finding}")
            
            p = doc.add_paragraph()
            
            if result.original_clause is not None:
                # Strikethrough original text
                run = p.add_run(result.original_clause.text + "\n")
                run.font.strike = True
                run.font.color.rgb = RGBColor(220, 38, 38)
            else:
                p.add_run("Clause not found — addition recommended\n").italic = True
                
            if result.suggested_redline:
                # Green bold suggested text
                run = p.add_run(result.suggested_redline)
                run.font.bold = True
                run.font.color.rgb = RGBColor(22, 101, 52)
                
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
